"""Build the NNDL submission register (Word) -- one table per Sir's format.

    python scripts/build_register.py [links.json]

links.json (optional) maps item key -> URL, plus "dataport"/"kaggle". Nothing is
invented: a missing link prints PASTE LINK HERE in red, a missing status stays as given.
"""
import json, sys
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent
links = json.loads(Path(sys.argv[1]).read_text()) if len(sys.argv) > 1 else {}

ME = dict(roll="CB.SC.U4CSE23134", name="Radha Krishna", group="1 (Team 1)",
          title="Silent Voice: Indian Sign Language to English translator (BiLSTM vs 1D CNN)")

# key, item, topic, due, status (honest, from the repo), file
ITEMS = [
    ("lab1", "Lab 1", "Lab 1 submission", "-", "Own file - add link", ""),
    ("lab2", "Lab 2 (15-Oct)", "Test and Evaluation (CDAO T&E framework) for Silent Voice", "15-Oct-2026 9:40 am", "Ready", "CB.SC.U4CSE23134.DOCX"),
    ("ezcnn", "Ez-CNN", "Ez-CNN", "-", "Own file - add link", ""),
    ("kmeans", "K-means with KNN", "K-means with KNN", "-", "Own file - add link", ""),
    ("act", "Activation functions", "Activation functions", "-", "Own file - add link", ""),
    ("r1", "Review 1", "Case study: BiLSTM vs 1D CNN for ISL recognition", "-", "Ready", "01_Review1_Case_Study_PPT.pptx"),
    ("r2", "Review 2", "Architecture, metrics, IEEE report, rubric deck", "-", "Ready", "04_Review2_PPT, 05_Review2_Rubric_PPT, 11_Review2_IEEE_Report"),
    ("r3", "Review 3 (15-Oct)", "Final review", "15-Oct-2026", "NOT prepared", ""),
]

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)

def put(cell, text, bold=False, red=False, size=9):
    cell.text = ""
    r = cell.paragraphs[0].add_run(text); r.bold = bold; r.font.size = Pt(size)
    if red: r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

def link(cell, key):
    u = links.get(key)
    put(cell, u if u else "PASTE LINK HERE", red=not u)

d = Document()
s = d.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = s.page_height, s.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, m, Cm(1.5))
d.styles["Normal"].font.name = "Calibri"

h = d.add_paragraph(); r = h.add_run("23CSE461 NNDL - Submission Link Register"); r.bold = True; r.font.size = Pt(16)
d.add_paragraph("Deadline: 15-Oct-2026, 4:00 PM (soft copy Word + hardcopy of status; also post the message in the WhatsApp group).").runs[0].font.size = Pt(9)

d.add_paragraph().add_run("1. Student and dataset details").bold = True
t = d.add_table(rows=2, cols=6); t.style = "Table Grid"
for i, x in enumerate(["Roll no", "Name", "Group no", "Title", "IEEE DataPort URL", "Other dataset URL (Kaggle etc.)"]):
    put(t.rows[0].cells[i], x, bold=True); shade(t.rows[0].cells[i], "D9E2F3")
put(t.rows[1].cells[0], ME["roll"]); put(t.rows[1].cells[1], ME["name"]); put(t.rows[1].cells[2], ME["group"]); put(t.rows[1].cells[3], ME["title"])
link(t.rows[1].cells[4], "dataport"); link(t.rows[1].cells[5], "kaggle")

d.add_paragraph()
d.add_paragraph().add_run("2. Submissions (single OneDrive folder; a heading and URL for each)").bold = True
hdr = ["Sl", "Roll no", "Name", "Heading", "Topic / title", "OneDrive URL", "Status", "Hardcopy"]
t = d.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
for i, x in enumerate(hdr):
    put(t.rows[0].cells[i], x, bold=True); shade(t.rows[0].cells[i], "D9E2F3")
for n, (key, item, topic, due, status, _) in enumerate(ITEMS, 1):
    c = t.add_row().cells
    put(c[0], str(n)); put(c[1], ME["roll"]); put(c[2], ME["name"]); put(c[3], item, bold=True); put(c[4], topic)
    link(c[5], key); put(c[6], status, red=status.startswith("NOT")); put(c[7], "To print / bring")

d.add_paragraph()
p = d.add_paragraph(); p.add_run("Note: ").bold = True
p.add_run("Each URL must open for the faculty (set the OneDrive link to 'Anyone with the link can view'). "
          "The IEEE DataPort URL is filled only after the dataset is published; the INCLUDE source data is "
          "Zenodo (CC BY-NC-SA 4.0) and is cited, not re-uploaded.")
for r in p.runs: r.font.size = Pt(9)

out = ROOT / "docs" / "NNDL_Submission_Link_Register_CB.SC.U4CSE23134.docx"
d.save(out); print("saved", out)
